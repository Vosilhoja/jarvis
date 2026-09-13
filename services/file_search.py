"""
Полнотекстовый поиск по СОДЕРЖИМОМУ файлов (не по имени — для поиска по имени
см. services/misc_tools.py::find_file_broad / search_files).

Реализовано осторожно, согласно AGENT_RULES.md: неограниченный обход реального
диска пользователя с чтением содержимого каждого файла может быть очень
медленным или зависнуть. Поэтому:
  - поиск ограничен стандартными пользовательскими папками (Рабочий стол,
    Документы, Загрузки) либо явно указанной пользователем директорией;
  - ограничены типы файлов (.txt, .md, .csv, .log, .docx, .pdf);
  - есть общий бюджет времени (TIME_BUDGET_SEC) и лимит числа файлов
    (MAX_FILES_SCANNED), проверяемые в цикле — как только исчерпаны, поиск
    прерывается с частичным результатом, а не зависает;
  - каждый файл ограничен по размеру (MAX_FILE_SIZE_BYTES) — гигантские файлы
    просто пропускаются;
  - PDF читается не полностью, а только первые MAX_PDF_PAGES страниц каждого
    файла, чтобы одна огромная книга не съела весь бюджет времени.
"""
from __future__ import annotations

import time
import logging
from pathlib import Path
from typing import Iterable, List, Optional, Tuple

logger = logging.getLogger("jarvis")

SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".log", ".docx", ".pdf"}
MAX_FILE_SIZE_BYTES = 15 * 1024 * 1024  # 15 МБ — больше не читаем
MAX_PDF_PAGES = 40
TIME_BUDGET_SEC = 25.0
MAX_FILES_SCANNED = 4000
MAX_MATCHES = 10
SNIPPET_RADIUS = 70  # символов вокруг найденного места


def _extract_text(path: Path) -> Optional[str]:
    """Достаёт текст из файла в зависимости от расширения. None — если не смогли/не поддерживается."""
    ext = path.suffix.lower()
    try:
        if ext in (".txt", ".md", ".csv", ".log"):
            for enc in ("utf-8", "cp1251", "cp866"):
                try:
                    return path.read_text(encoding=enc, errors="ignore")
                except Exception:
                    continue
            return None

        if ext == ".docx":
            try:
                import docx  # python-docx
            except ImportError:
                logger.warning("file_search: пакет python-docx не установлен — .docx пропускаются")
                return None
            doc = docx.Document(str(path))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        parts.append(cell.text)
            return "\n".join(parts)

        if ext == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                logger.warning("file_search: пакет pypdf не установлен — .pdf пропускаются")
                return None
            reader = PdfReader(str(path))
            pages = reader.pages[:MAX_PDF_PAGES]
            return "\n".join((p.extract_text() or "") for p in pages)
    except Exception as e:
        logger.debug(f"file_search: не удалось прочитать {path}: {e}")
        return None
    return None


def _make_snippet(text: str, idx: int, query_len: int) -> str:
    start = max(0, idx - SNIPPET_RADIUS)
    end = min(len(text), idx + query_len + SNIPPET_RADIUS)
    snippet = text[start:end].replace("\n", " ").strip()
    prefix = "…" if start > 0 else ""
    suffix = "…" if end < len(text) else ""
    return f"{prefix}{snippet}{suffix}"


def _default_roots() -> List[Path]:
    home = Path.home()
    return [home / "Desktop", home / "Documents", home / "Downloads"]


def search_file_content(
    query: str,
    root: Optional[str] = None,
    extensions: Optional[Iterable[str]] = None,
) -> str:
    """
    Ищет строку `query` внутри содержимого файлов. Возвращает готовое
    Markdown-сообщение для Telegram (список путей + короткий контекст находки).
    """
    query = (query or "").strip()
    if not query:
        return "⚠️ Пустой поисковый запрос."

    if root:
        from core.execution.common import resolve_path_aliases
        roots = [resolve_path_aliases(root)]
    else:
        roots = _default_roots()

    allowed_exts = {e.lower() if e.startswith(".") else f".{e.lower()}" for e in extensions} if extensions else SUPPORTED_EXTENSIONS

    query_lower = query.lower()
    matches: List[Tuple[Path, str]] = []
    files_scanned = 0
    start = time.time()
    timed_out = False

    for base in roots:
        if not base.exists():
            continue
        try:
            for p in base.rglob("*"):
                if time.time() - start > TIME_BUDGET_SEC:
                    timed_out = True
                    break
                if files_scanned > MAX_FILES_SCANNED:
                    timed_out = True
                    break
                if not p.is_file() or p.suffix.lower() not in allowed_exts:
                    continue
                try:
                    if p.stat().st_size > MAX_FILE_SIZE_BYTES:
                        continue
                except Exception:
                    continue

                files_scanned += 1
                text = _extract_text(p)
                if not text:
                    continue
                idx = text.lower().find(query_lower)
                if idx == -1:
                    continue
                snippet = _make_snippet(text, idx, len(query))
                matches.append((p, snippet))
                if len(matches) >= MAX_MATCHES:
                    break
        except Exception as e:
            logger.debug(f"file_search: ошибка обхода {base}: {e}")
        if len(matches) >= MAX_MATCHES or timed_out:
            break

    if not matches:
        suffix = " (поиск прерван по таймауту, результат может быть неполным)" if timed_out else ""
        return f"🔍 По запросу «{query}» ничего не найдено в содержимом файлов{suffix}."

    lines = [f"🔍 *Найдено в содержимом файлов* (по запросу «{query}»):\n"]
    for p, snippet in matches:
        lines.append(f"📄 `{p}`\n   ...{snippet}...\n")
    lines.append(f"\n_Просканировано файлов: {files_scanned}. Типы: {', '.join(sorted(allowed_exts))}._")
    if timed_out:
        lines.append("_⏳ Поиск остановлен по лимиту времени/файлов — результат может быть неполным._")
    return "\n".join(lines)
